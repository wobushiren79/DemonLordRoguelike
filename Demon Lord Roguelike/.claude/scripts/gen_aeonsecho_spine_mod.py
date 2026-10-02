# -*- coding: utf-8 -*-
"""
AeonsEchoSpine Mod 数据生成器（Demon Lord Roguelike）

两段式流水线（Excel 为唯一真实源，与主项目配置表惯例一致）：
  scan   扫描 MOD 项目 Assets/ModResource/Spine/AeonsEcho 资源套装 → 重建道具 Excel（合并保留语言 Excel 的人工改名）
  export 读道具/语言 Excel → 导出 Mod JsonText（ItemsInfo.txt + Language_ItemsInfo_*.txt ×12）
  all    scan + export（默认）
  migrate 道具表 other_data 旧位置段格式 → 键值格式（一次性迁移，幂等，数值原样保留）

other_data 键值格式（& 拆项、每项首个 : 拆键值，同主项目 attack_mode other_data 规约，缺省键省略）：
  show_res:1101_Chess_SkeletonData&ui_show_res:1101_Avator_lv1_SkeletonData&ui_show_data:0.1919;0,0&show_data:3.75;0,-120&world_data:1;0,0&idle_anim:idle_battle&ui_show_idle_anim:00_Idle
  show_res     默认展示形象资源名（可空：套装无 Chess 时省略=仅详情UI幻化，世界/小卡回落原生物形象；show=游戏默认展示，即原 chess 概念）
  ui_show_res  ui_show_spine 高清展示资源名（可空，无 ui_show 变体时省略，ui_show=详情UI高清展示，即原 avator 概念；Avator/Secretary/Elf/AVG 同规则）
  ui_show_data 详情UI尺寸「scale;x,y」（有 ui_show 变体时带，格式同 CreatureModelBean.ui_data_b）
  show_data    默认展示小卡UI尺寸「scale;x,y」（有 Chess 时带，格式同 CreatureModelBean.ui_data_s，即原 ui_chess_data）
  world_data   世界显示尺寸/偏移「scale;x,y」（可空，x=横向偏移,y=竖向抬升，战斗/基地等世界空间 SkeletonAnimation 消费；
               手调数据无骨架校准来源——主项目幻化药测试面板场景列表可编辑，scan 全量重建时按 id 保留，不会被校准覆盖）
  idle_anim    show 骨架的替代待机动画名（可空：骨架动画列表命中主项目标准待机候选[excel_spine_animation_state id=10001 的 res 字段,当前 idle,wait,idle1,wait1,stand]时省略=走框架候选解析；
               无标准候选时取首个小写含 idle 的动画名[保留大小写,Spine 需精确名]；完全没有含 idle 动画则不生成该键=保持现状静态+警告）
  ui_show_idle_anim ui_show 骨架的替代待机动画名（可空，规则同 idle_anim，检测对象为 ui_show 变体骨架）

套装组合规则：Chess×ui_show 笛卡尔积；只有 Chess → 每个 Chess 单独出道具（只设 show 段）；
  只有 ui_show（Avator/Secretary/Elf/AVG）→ 每个 ui_show 变体单独出道具（只设 ui_show 段，详情UI幻化）；
  两类都没有 → 跳过该套装。
  混合套装序号排序 = ui_show 前缀组（Avator→Secretary→Elf→AVG）→ Chess → 组内变体自然序
  （单组套装与旧版 Chess×Avator 顺序完全一致；新前缀组的组合整体排后，旧套装 id→组合映射不漂移）

Excel 位置（MOD 项目，可用 Excel/WPS 直接打开查看/调整）：
  Assets/Data/Excel/excel_mod_items_info[Mod道具信息].xlsx   - 道具配置（3 行表头：列名/类型/说明，同主项目布局）
  Assets/Data/Excel/excel_mod_language[Mod多语言].xlsx       - 道具名多语言（id + content_{12语言}，改名在这里）

注意：
  - scan 会全量重建道具表数据行（覆盖前先备份到 MOD项目/ExcelBackup/，不在 Assets 内），但 **show_data/ui_show_data/world_data 手调值按 id 保留**（新增资源仍按骨架校准默认值；--reset-layout 可强制全部重算）；改了道具表参数后请只跑 export
  - 语言表按 id 合并保留人工内容：新增道具补默认名，删除的道具行自动清理
  - 规则详见主项目 .claude/skills/aeonsecho-spine-mod/SKILL.md

用法（一律走主项目 .claude/scripts/run-python.ps1 包装）：
  powershell.exe -NoProfile -ExecutionPolicy Bypass -File ".claude/scripts/run-python.ps1" `
      ".claude/scripts/gen_aeonsecho_spine_mod.py" all --mod-project "<MOD项目根>"
  # 构建完成后部署到主项目：
  ... export --mod-project "<MOD项目根>" --deploy-main "<主项目根>"
"""
import argparse
import json
import shutil
import sys
from pathlib import Path

import openpyxl

MOD_NAME = "AeonsEchoSpine"
SOURCE_REL = "Assets/ModResource/Spine/AeonsEcho"
ICON_RES = "Item_TransformPotion_1"  # 复用主游戏内置幻化药图标
ITEM_TYPE_TRANSFORM_POTION = 18
LANGUAGES = ["cn", "en", "jp", "kr", "tw", "de", "fr", "ru", "es", "br", "pl", "tr"]

# 道具表：列名/类型/说明（布局与主项目 excel_items_info 一致，name 带 [language] 标记约定）
ITEM_SHEET = "ItemsInfo"
ITEM_EXCEL_REL = "Assets/Data/Excel/excel_mod_items_info[Mod道具信息].xlsx"
# ui_show 高清展示变体前缀（详情UI形象，走 ui_show_res/ui_show_data 键；Avator/Secretary/Elf/AVG 同规则）
UI_SHOW_PREFIXES = ("Avator", "Secretary", "Elf", "AVG")
ITEM_COLUMNS = [
    ("id", "long", "道具ID(key)=18+套装号+2位序号"),
    ("item_type", "int", "道具类型(18=幻化药)"),
    ("item_weapon_type", "int", "武器类型(仅武器有效)"),
    ("num_max", "int", "堆叠上限(1=不堆叠)"),
    ("creature_model_id", "long", "生物模组信息ID(0=通用)"),
    ("creature_model_info_id", "long", "生物模组详细信息ID"),
    ("icon_res", "string", "图标资源"),
    ("icon_rotate_z", "float", "图标旋转"),
    ("attack_mode_data", "string", "攻击模式数据(幻化药不用)"),
    ("other_data", "string", "形象键值串:show段(show_res:X&show_data:scale;x,y)与ui_show段(ui_show_res:X&ui_show_data:scale;x,y)至少一段,可选world_data:scale;x,y、idle_anim/ui_show_idle_anim:动画名(无标准idle时的替代待机动画)(&拆项,:拆键值,缺省键省略)"),
    ("name[language]", "long", "道具名textId(=道具id,文本在excel_mod_language)"),
    ("remark", "string", "备注"),
    ("reward_rarity", "string", "奖励稀有度白名单(空=全适配;消耗品不进装备池)"),
    ("source", "string", "道具来源(逗号分隔枚举值,空=默认来源) 1=征服模式奖励"),
]

# 语言表：id + content_{lang}（与主项目 excel_language 的 sheet 布局一致）
LANG_SHEET = "ItemsInfo"
LANG_EXCEL_REL = "Assets/Data/Excel/excel_mod_language[Mod多语言].xlsx"


def natural_key(name: str):
    """自然排序键（数字段按数值）：Chess < Chess_s01 < Chess_s02，Avator_lv1 < Avator_lv1a < Avator_lv2"""
    import re
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", name)]


def ui_show_group_index(name: str) -> int:
    """ui_show 变体的前缀组序号（Avator→Secretary→Elf→AVG）；用于混合套装组合排序——
    新前缀组的组合整体排在旧组之后，保证旧套装已出道具的 id→组合映射不漂移"""
    for idx, prefix in enumerate(UI_SHOW_PREFIXES):
        if name.startswith(prefix):
            return idx
    return len(UI_SHOW_PREFIXES)


def default_name(lang: str, set_id: str, seq: int) -> str:
    """道具名默认值：cn/tw 中文，其余语言英文兜底"""
    if lang == "cn":
        return f"幻化药·回响{set_id}-{seq:02d}"
    if lang == "tw":
        return f"幻化藥·迴響{set_id}-{seq:02d}"
    return f"Echo Potion {set_id}-{seq:02d}"


def get_spine_height(json_path: Path) -> float:
    """读 spine json 的骨架包围盒高度（用于校准详情UI缩放）"""
    try:
        data = json.loads(json_path.read_text(encoding="utf-8"))
        return float(data["skeleton"]["height"])
    except Exception:
        return 0.0


# 标准待机动画候选的兜底值（=主项目 excel_spine_animation_state id=10001 的 res 字段；读不到配置时用）
STD_IDLE_FALLBACK = ["idle", "wait", "idle1", "wait1", "stand"]


def load_std_idle_candidates() -> list:
    """读主项目 SpineAnimationState.txt 的 Idle(id=10001) 候选名列表（配置表为唯一真实源,不硬编码；
    按脚本位置推导主项目根=.claude/scripts/ 向上两级）；读失败回退 STD_IDLE_FALLBACK"""
    try:
        main_root = Path(__file__).resolve().parents[2]
        rows = json.loads((main_root / "Assets/Resources/JsonText/SpineAnimationState.txt").read_text(encoding="utf-8"))
        for row in rows:
            if int(row.get("id", 0)) == 10001:
                candidates = [s.strip() for s in str(row.get("res", "")).split(",") if s.strip()]
                if candidates:
                    return candidates
    except Exception:
        pass
    return STD_IDLE_FALLBACK


def get_spine_anims(json_path: Path) -> list:
    """读 spine json 的动画名列表（4.x 格式 animations 为 dict,取 keys 保持声明顺序）；异常返回空列表"""
    try:
        data = json.loads(json_path.read_text(encoding="utf-8"))
        anims = data.get("animations")
        if isinstance(anims, dict):
            return list(anims.keys())
    except Exception:
        pass
    return []


def pick_idle_anim(anims: list, std_candidates: list) -> str:
    """检测骨架动画列表的待机动画（idle 动画替代规则,新 Mod 生成脚本同样继承）：
    命中标准候选(大小写不敏感全等)→返回""(不需替代,运行时走框架候选解析)；
    否则取首个小写含 idle 的动画原始名(保留大小写,Spine SetAnimation 需精确名)；
    完全没有含 idle 动画→返回""(不生成 idle_anim 键,保持现状静态显示)"""
    lowers = {a.lower() for a in anims}
    for c in std_candidates:
        if c.lower() in lowers:
            return ""
    for a in anims:
        if "idle" in a.lower():
            return a
    return ""


def backup_excel(excel_path: Path, mod_project: Path):
    """覆盖 Excel 前备份到 MOD项目/ExcelBackup/（Assets 之外，不被导出工具扫描、不产生 .meta）；
    滚动复用 .bak.1~.bak.3（1=最新）——移位覆盖旧文件不新增，目录里永远只有最近 3 份，并顺带清掉旧版时间戳命名备份"""
    if not excel_path.exists():
        return
    backup_dir = mod_project / "ExcelBackup"
    backup_dir.mkdir(exist_ok=True)
    name = excel_path.name
    # 移位覆盖：bak.3←bak.2、bak.2←bak.1（replace 直接覆盖复用旧文件，不产生新文件）
    for i in range(3, 1, -1):
        newer = backup_dir / f"{name}.bak.{i - 1}"
        if newer.exists():
            newer.replace(backup_dir / f"{name}.bak.{i}")
    shutil.copy2(excel_path, backup_dir / f"{name}.bak.1")
    # 清理非滚动命名的旧备份（时间戳版等），保证永远都只有 3 份
    keep = {f"{name}.bak.{i}" for i in range(1, 4)}
    for old in backup_dir.iterdir():
        if old.is_file() and old.name.startswith(f"{name}.bak.") and old.name not in keep:
            old.unlink()


def build_other_data(show_res: str, ui_show_res: str = "", ui_show_data: str = "", show_data: str = "", world_data: str = "", idle_anim: str = "", ui_show_idle_anim: str = "") -> str:
    """拼 other_data 键值串：& 拆项、每项首个 : 拆键值（同主项目 attack_mode other_data 规约），缺省键省略；
    show_res 可空——仅详情UI幻化道具（套装无 Chess，只有 ui_show 变体）只有 ui_show 段；
    idle_anim/ui_show_idle_anim=show/ui_show 骨架无标准待机动画时的替代动画名（idle 动画替代规则自动检测,命中候选/无 idle 均省略）"""
    segs = []
    if show_res:
        segs.append(f"show_res:{show_res}")
    if ui_show_res:
        segs.append(f"ui_show_res:{ui_show_res}")
    if ui_show_data:
        segs.append(f"ui_show_data:{ui_show_data}")
    if show_data:
        segs.append(f"show_data:{show_data}")
    if world_data:
        segs.append(f"world_data:{world_data}")
    if idle_anim:
        segs.append(f"idle_anim:{idle_anim}")
    if ui_show_idle_anim:
        segs.append(f"ui_show_idle_anim:{ui_show_idle_anim}")
    return "&".join(segs)


def migrate_other_data(old: str) -> str:
    """旧位置段格式（chessRes[,avatorRes|uiShowData|uiChessData]）→ 键值格式，数值原样保留；含 : 的视为新格式原样返回（幂等）"""
    s = (old or "").strip()
    if not s or ":" in s:
        return s
    parts = s.split("|")
    res_split = parts[0].split(",", 1)
    show_res = res_split[0].strip()
    ui_show_res = res_split[1].strip() if len(res_split) > 1 else ""
    ui_show = parts[1].strip() if len(parts) > 1 else ""
    show_data = parts[2].strip() if len(parts) > 2 else ""
    return build_other_data(show_res, ui_show_res, ui_show, show_data)


def migrate_item_excel(excel_path: Path, mod_project: Path):
    """道具表 other_data 旧格式→键值格式一次性迁移；有改动才备份写回"""
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[ITEM_SHEET]
    header = [c.value for c in ws[1]]
    if "other_data" not in header:
        wb.close()
        print("[错误] 道具表缺少 other_data 列")
        sys.exit(1)
    col = header.index("other_data") + 1
    migrated = skipped = 0
    for row in range(4, ws.max_row + 1):
        if ws.cell(row=row, column=1).value is None:
            continue
        old = ws.cell(row=row, column=col).value
        old_s = "" if old is None else str(old)
        new = migrate_other_data(old_s)
        if new != old_s:
            ws.cell(row=row, column=col, value=new)
            migrated += 1
        else:
            skipped += 1
    if migrated > 0:
        backup_excel(excel_path, mod_project)
        wb.save(excel_path)
    wb.close()
    print(f"[migrate] other_data 旧格式→键值格式：转换 {migrated} 行，已是新格式/空 {skipped} 行 → {excel_path}")


# scan 重建时按 id 保留的布局键（卡片位置/大小手调值）：小卡 show_data、详情UI ui_show_data、世界 world_data
PRESERVED_LAYOUT_KEYS = ("show_data", "ui_show_data", "world_data")


def read_preserved_layout_data(excel_path: Path) -> dict:
    """读现有道具表提取各行 other_data 的布局三键 → {道具id: {键: 值}}；
    scan 全量重建时默认按 id 保留测试面板手调的卡片位置/大小（同 id=同资源,骨架未变,手调值仍有效；
    新增资源无保留值按骨架校准默认值；--reset-layout 可强制全部重算；id 已失效的行自然丢弃）"""
    preserved = {}
    if not excel_path.exists():
        return preserved
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    if ITEM_SHEET not in wb.sheetnames:
        wb.close()
        return preserved
    ws = wb[ITEM_SHEET]
    header = [c.value for c in ws[1]]
    if "id" not in header or "other_data" not in header:
        wb.close()
        return preserved
    col_id = header.index("id")
    col_other = header.index("other_data")
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None or row[col_other] is None:
            continue
        data = {}
        for seg in str(row[col_other]).split("&"):
            for key in PRESERVED_LAYOUT_KEYS:
                if seg.startswith(key + ":"):
                    data[key] = seg[len(key) + 1:].strip()
                    break
        if data:
            preserved[int(row[col_id])] = data
    wb.close()
    return preserved


def detect_idle_anim(json_path: Path, std_candidates: list, idle_stats: dict, warnings: list, tag: str) -> str:
    """对单个骨架执行 idle 动画替代规则检测并累计统计；tag=警告里的资源标识。
    返回替代待机动画名（命中标准候选/无 idle 均返回 ""）"""
    anims = get_spine_anims(json_path)
    idle_anim = pick_idle_anim(anims, std_candidates)
    if not anims:
        idle_stats["none"] += 1
        warnings.append(f"{tag} 动画列表读取失败,未检测 idle 动画")
    elif idle_anim:
        idle_stats["replaced"] += 1
    elif not any(c.lower() in {a.lower() for a in anims} for c in std_candidates):
        idle_stats["none"] += 1
        warnings.append(f"{tag} 无任何 idle 动画(将静态显示)")
    else:
        idle_stats["std"] += 1
    return idle_anim


def compute_items(source_dir: Path, ui_scale_k: float, ui_pos_y: float, ui_chess_scale_k: float, ui_chess_pos_y: float, preserved_layout: dict = None, reset_layout: bool = False, std_idle_candidates: list = None):
    """扫描资源套装计算道具行，返回 (道具行列表, 跳过套装清单, 警告清单, idle统计dict)；道具行 key 为干净字段名；
    preserved_layout=scan 前从旧道具表读出的布局三键手调值（按 id 保留）；reset_layout=True 时忽略保留值全部按骨架重算；
    std_idle_candidates=主项目标准待机动画候选名列表（idle 动画替代规则检测用,None 时跳过检测）"""
    preserved_layout = {} if reset_layout else (preserved_layout or {})
    # idle 动画检测统计（show/ui_show 两段合计）：std=命中标准候选, replaced=替代动画, none=完全无 idle
    idle_stats = {"std": 0, "replaced": 0, "none": 0}
    sets = []
    skipped = []
    for set_dir in sorted(source_dir.iterdir(), key=lambda p: natural_key(p.name)):
        if not set_dir.is_dir() or not set_dir.name.isdigit():
            continue
        chess_list = sorted(
            [d.name for d in set_dir.iterdir() if d.is_dir() and d.name.startswith("Chess")],
            key=natural_key)
        ui_show_list = sorted(
            [d.name for d in set_dir.iterdir() if d.is_dir() and d.name.startswith(UI_SHOW_PREFIXES)],
            key=natural_key)
        if not chess_list and not ui_show_list:
            skipped.append(set_dir.name)
            continue
        sets.append((set_dir.name, chess_list, ui_show_list))

    items = []
    used_ids = set()
    warnings = []
    for set_id, chess_list, ui_show_list in sets:
        # 缺 SkeletonData 资产的 Chess 变体先剔除（不占序号）
        valid_chess = []
        for chess in chess_list:
            if (source_dir / set_id / chess / f"{set_id}_{chess}_SkeletonData.asset").exists():
                valid_chess.append(chess)
            else:
                warnings.append(f"{set_id}/{chess}: 缺 SkeletonData 资产，跳过该变体")
        # 组合清单：Chess×ui_show 笛卡尔积；只有一类时该类每个变体单独出 1 个道具
        # （仅 Chess→只设 show 段；仅 ui_show→只设 ui_show 段=详情UI幻化，世界/小卡回落原生物形象）
        if valid_chess and ui_show_list:
            # ui_show 前缀组最外层（新组的组合整体排后，旧套装 id→组合映射不漂移），组内保持 Chess×变体(外层) 旧顺序
            combos = []
            for idx in range(len(UI_SHOW_PREFIXES) + 1):
                group = [u for u in ui_show_list if ui_show_group_index(u) == idx]
                combos.extend((c, u) for c in valid_chess for u in group)
        elif valid_chess:
            combos = [(c, None) for c in valid_chess]
        else:
            combos = [(None, u) for u in ui_show_list]

        seq = 0
        for chess, uishow in combos:
            # 仅 ui_show 组合资源缺失→道具无有效内容，跳过不占序号
            if chess is None and not (source_dir / set_id / uishow / f"{set_id}_{uishow}_SkeletonData.asset").exists():
                warnings.append(f"{set_id}/{uishow}: 缺 SkeletonData 资产，跳过该变体")
                continue
            seq += 1
            if seq > 99:
                warnings.append(f"{set_id}: 组合数超过99，ID序号溢出，后续组合未生成")
                break
            self_id = int(f"{ITEM_TYPE_TRANSFORM_POTION}{set_id}{seq:02d}")
            if self_id in used_ids:
                warnings.append(f"ID冲突: {self_id} ({set_id} 第{seq}个组合)")
                continue
            used_ids.add(self_id)
            # 布局三键手调值按 id 保留(同 id=同资源,骨架未变,手调值仍有效; --reset-layout 时 preserved_layout 已置空=全部重算)
            preserved = preserved_layout.get(self_id, {})
            # world_data 手调值按 id 保留(无校准来源, 默认不生成该键)
            world_data = preserved.get("world_data", "")

            show_res = chess_ui_data = ui_show_res = ui_show_data = idle_anim = ui_show_idle_anim = ""
            chess_tag = uishow_tag = ""
            if chess is not None:
                # show 段：默认展示形象 + 小卡尺寸（有保留值用手调值，否则按 Chess 骨架高校准 scale=K2/骨架高）
                show_res = f"{set_id}_{chess}_SkeletonData"
                chess_ui_data = preserved.get("show_data", "")
                if not chess_ui_data:
                    chess_height = get_spine_height(source_dir / set_id / chess / f"{set_id}_{chess}.json")
                    chess_ui_scale = round(ui_chess_scale_k / chess_height, 4) if chess_height > 0 else 3.75
                    chess_ui_data = f"{chess_ui_scale};0,{int(ui_chess_pos_y)}"
                # idle 动画替代规则(show 段):Chess 骨架无标准待机动画时取首个含 idle 的动画名写入 idle_anim 键
                if std_idle_candidates is not None:
                    idle_anim = detect_idle_anim(source_dir / set_id / chess / f"{set_id}_{chess}.json", std_idle_candidates, idle_stats, warnings, f"{set_id}/{chess}(show段):")
                chess_tag = chess
            if uishow is not None:
                # ui_show 段：详情UI高清形象 + 详情UI尺寸（有保留值用手调值，否则按 ui_show 骨架高校准 scale=K/骨架高）
                if not (source_dir / set_id / uishow / f"{set_id}_{uishow}_SkeletonData.asset").exists():
                    # 混合组合中 ui_show 资产缺失：该组合退化为仅 show（序号已消耗）
                    warnings.append(f"{set_id}/{uishow}: 缺 SkeletonData 资产，该组合按无 ui_show 处理")
                else:
                    ui_show_res = f"{set_id}_{uishow}_SkeletonData"
                    ui_show_data = preserved.get("ui_show_data", "")
                    if not ui_show_data:
                        height = get_spine_height(source_dir / set_id / uishow / f"{set_id}_{uishow}.json")
                        ui_scale = round(ui_scale_k / height, 4) if height > 0 else 0.12
                        ui_show_data = f"{ui_scale};0,{int(ui_pos_y)}"
                    # idle 动画替代规则(ui_show 段):ui_show 骨架无标准待机动画时取首个含 idle 的动画名写入 ui_show_idle_anim 键
                    if std_idle_candidates is not None:
                        ui_show_idle_anim = detect_idle_anim(source_dir / set_id / uishow / f"{set_id}_{uishow}.json", std_idle_candidates, idle_stats, warnings, f"{set_id}/{uishow}(ui_show段):")
                    uishow_tag = f"×{uishow}" if chess_tag else uishow
            other_data = build_other_data(show_res, ui_show_res, ui_show_data, chess_ui_data, world_data=world_data, idle_anim=idle_anim, ui_show_idle_anim=ui_show_idle_anim)

            items.append({
                "id": self_id,
                "item_type": ITEM_TYPE_TRANSFORM_POTION,
                "item_weapon_type": 0,
                "num_max": 1,
                "creature_model_id": 0,
                "creature_model_info_id": 0,
                "icon_res": ICON_RES,
                "icon_rotate_z": 0.0,
                "attack_mode_data": "",
                "other_data": other_data,
                "name": self_id,
                "remark": f"回响幻化药(AeonsEcho {set_id} {chess_tag}{uishow_tag})",
                "reward_rarity": "",
                # 来源=征服模式奖励(ItemSourceEnum.ConquerReward)：征服通关领奖的魔晶位之一随机替换为幻化药
                "source": "1",
                "_set_id": set_id,  # 内部字段：供语言默认名使用，不进 Excel/JSON
                "_seq": seq,
            })
    return items, skipped, warnings, idle_stats


def write_item_excel(excel_path: Path, items, mod_project: Path):
    """全量重建道具表（3 行表头：列名/类型/说明；数据从第 4 行起）"""
    backup_excel(excel_path, mod_project)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ITEM_SHEET
    for col, (name, type_name, comment) in enumerate(ITEM_COLUMNS, start=1):
        ws.cell(row=1, column=col, value=name)
        ws.cell(row=2, column=col, value=type_name)
        ws.cell(row=3, column=col, value=comment)
    for row, item in enumerate(items, start=4):
        for col, (name, _, _) in enumerate(ITEM_COLUMNS, start=1):
            key = name.replace("[language]", "")
            ws.cell(row=row, column=col, value=item[key])
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)


def merge_language_excel(excel_path: Path, items, mod_project: Path):
    """语言表按 id 合并：保留人工改名，新道具补默认名，失效 id 清理"""
    existing = {}  # id -> {lang: content}
    if excel_path.exists():
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        if LANG_SHEET in wb.sheetnames:
            ws = wb[LANG_SHEET]
            header = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=4, values_only=True):
                if row[0] is None:
                    continue
                row_id = int(row[0])
                existing[row_id] = {}
                for i, col_name in enumerate(header):
                    if col_name and col_name.startswith("content_") and i < len(row):
                        lang = col_name.replace("content_", "")
                        if row[i] is not None and str(row[i]).strip() != "":
                            existing[row_id][lang] = str(row[i])
        wb.close()

    backup_excel(excel_path, mod_project)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = LANG_SHEET
    ws.cell(row=1, column=1, value="id")
    ws.cell(row=2, column=1, value="long")
    ws.cell(row=3, column=1, value="道具名textId(key)=道具id")
    for col, lang in enumerate(LANGUAGES, start=2):
        ws.cell(row=1, column=col, value=f"content_{lang}")
        ws.cell(row=2, column=col, value="string")
        ws.cell(row=3, column=col, value="道具名(空=英文兜底)")
    kept = 0
    for row, item in enumerate(items, start=4):
        self_id = item["id"]
        ws.cell(row=row, column=1, value=self_id)
        old = existing.get(self_id, {})
        for col, lang in enumerate(LANGUAGES, start=2):
            content = old.get(lang) or default_name(lang, item["_set_id"], item["_seq"])
            if old.get(lang):
                kept += 1
            ws.cell(row=row, column=col, value=content)
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)
    return kept


def read_item_excel(excel_path: Path):
    """读道具表 → JSON 行（key 剥离 [language] 标记，按类型还原数值）"""
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb[ITEM_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = rows[0]
    types = rows[1]
    result = []
    for row in rows[3:]:
        if row[0] is None:
            continue
        item = {}
        for i, col_name in enumerate(header):
            if col_name is None:
                continue
            key = str(col_name).replace("[language]", "")
            type_name = str(types[i]) if types[i] else "string"
            value = row[i] if i < len(row) else None
            if type_name in ("long", "int"):
                item[key] = int(value) if value is not None and str(value).strip() != "" else 0
            elif type_name == "float":
                item[key] = float(value) if value is not None and str(value).strip() != "" else 0.0
            else:
                item[key] = "" if value is None else str(value)
        result.append(item)
    result.sort(key=lambda x: x["id"])
    return result


def read_language_excel(excel_path: Path):
    """读语言表 → {lang: [{id, content}]}（空内容按 cn→en 顺序兜底）"""
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb[LANG_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(c) if c else "" for c in rows[0]]
    per_lang = {lang: {} for lang in LANGUAGES}
    for row in rows[3:]:
        if row[0] is None:
            continue
        row_id = int(row[0])
        contents = {}
        for i, col_name in enumerate(header):
            if col_name.startswith("content_") and i < len(row) and row[i] is not None and str(row[i]).strip() != "":
                contents[col_name.replace("content_", "")] = str(row[i])
        for lang in LANGUAGES:
            # 兜底链：本语言 → cn → en → id
            content = contents.get(lang) or contents.get("cn") or contents.get("en") or str(row_id)
            per_lang[lang][row_id] = content
    return per_lang


def main():
    parser = argparse.ArgumentParser(description="AeonsEchoSpine Mod 数据生成器（Excel→JsonText 两段式）")
    parser.add_argument("command", nargs="?", default="all", choices=["scan", "export", "all", "migrate"],
                        help="scan=资源→Excel；export=Excel→JsonText；all=两者（默认）；migrate=道具表 other_data 旧格式→键值格式")
    parser.add_argument("--mod-project", required=True, help="MOD 项目根目录")
    parser.add_argument("--deploy-main", default="", help="可选：主项目根目录，构建完成后把 Mod 整个目录部署过去")
    parser.add_argument("--ui-scale-k", type=float, default=645.0, help="详情UI缩放校准常数 K：scale=K/Avator骨架高（默认645，=原430基准×1.5）")
    parser.add_argument("--ui-pos-y", type=float, default=0.0, help="详情UI默认Y偏移（默认0）")
    parser.add_argument("--ui-chess-scale-k", type=float, default=3159.0, help="小卡UI缩放校准常数 K2：scale=K2/Chess骨架高（默认3159=842.4×2.5×1.5，2.5=主游戏人形骨架 ui_data_s 基准，1.5=放大倍率）")
    parser.add_argument("--ui-chess-pos-y", type=float, default=-120.0, help="小卡UI默认Y偏移（默认-120）")
    parser.add_argument("--reset-layout", action="store_true",
                        help="重建时重新设置卡片位置/大小（show_data/ui_show_data/world_data 全部按骨架重算）；默认关闭=按 id 保留已有手调值，仅新资源计算默认值")
    args = parser.parse_args()

    mod_project = Path(args.mod_project)
    source_dir = mod_project / SOURCE_REL
    item_excel = mod_project / ITEM_EXCEL_REL
    lang_excel = mod_project / LANG_EXCEL_REL

    if args.command == "migrate":
        if not item_excel.exists():
            print("[错误] 道具 Excel 不存在，请先运行 scan")
            sys.exit(1)
        migrate_item_excel(item_excel, mod_project)

    if args.command in ("scan", "all"):
        if not source_dir.is_dir():
            print(f"[错误] 资源目录不存在: {source_dir}")
            sys.exit(1)
        preserved_layout = read_preserved_layout_data(item_excel)
        std_idle_candidates = load_std_idle_candidates()
        items, skipped, warnings, idle_stats = compute_items(source_dir, args.ui_scale_k, args.ui_pos_y, args.ui_chess_scale_k, args.ui_chess_pos_y, preserved_layout, reset_layout=args.reset_layout, std_idle_candidates=std_idle_candidates)
        write_item_excel(item_excel, items, mod_project)
        kept = merge_language_excel(lang_excel, items, mod_project)
        print(f"[scan] 重建道具表 {len(items)} 行 → {item_excel}")
        print(f"[scan] 合并语言表（保留人工内容 {kept} 格）→ {lang_excel}")
        print(f"[scan] idle 动画检测：命中标准候选 {idle_stats['std']}，替代动画 {idle_stats['replaced']}，无 idle {idle_stats['none']}（标准候选={','.join(std_idle_candidates)}）")
        if args.reset_layout:
            print("[scan] --reset-layout：布局三键(show_data/ui_show_data/world_data)全部按骨架重算")
        elif preserved_layout:
            print(f"[scan] 按 id 保留布局手调值 {len(preserved_layout)} 行（show_data/ui_show_data/world_data；--reset-layout 可强制重算）")
        if skipped:
            print(f"[scan] 跳过套装（无 Chess/Avator/Secretary/Elf/AVG）: {', '.join(skipped)}")
        for w in warnings:
            print(f"[scan] [警告] {w}")

    if args.command in ("export", "all"):
        if not item_excel.exists() or not lang_excel.exists():
            print("[错误] Excel 不存在，请先运行 scan")
            sys.exit(1)
        items = read_item_excel(item_excel)
        per_lang = read_language_excel(lang_excel)

        out_dir = mod_project / "Mods" / MOD_NAME / "JsonText"
        out_dir.mkdir(parents=True, exist_ok=True)
        dump = lambda rows: json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
        (out_dir / "ItemsInfo.txt").write_text(dump(items), encoding="utf-8")
        for lang in LANGUAGES:
            rows = [{"id": k, "content": v} for k, v in sorted(per_lang[lang].items())]
            (out_dir / f"Language_ItemsInfo_{lang}.txt").write_text(dump(rows), encoding="utf-8")
        print(f"[export] 导出 {len(items)} 个道具 → {out_dir}")
        print(f"[export] ItemsInfo.txt + Language_ItemsInfo_*.txt ×{len(LANGUAGES)}")

    # 部署：构建完成后把 Mod 整个目录拷贝到主项目 Mods 下
    if args.deploy_main:
        mod_out = mod_project / "Mods" / MOD_NAME
        if not (mod_out / "catalog.bin").exists():
            print(f"[错误] 未找到构建产物 {mod_out}/catalog.bin —— 请先在 MOD 项目 Unity 里执行「工具/Mod/AeonsEchoSpine/一键构建」")
            sys.exit(1)
        target = Path(args.deploy_main) / "Mods" / MOD_NAME
        if target.exists():
            shutil.rmtree(target)
        shutil.copytree(mod_out, target)
        print(f"[deploy] 已部署 → {target}")
        print("下一步：主项目 Unity 里用 LauncherTest 启动（内含 InitializeAllModsSync），吃幻化药验证形象/详情UI")


if __name__ == "__main__":
    main()
