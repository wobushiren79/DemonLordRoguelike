# -*- coding: utf-8 -*-
"""
BrownDustSpine Mod 数据生成器（Demon Lord Roguelike）

两段式流水线（Excel 为唯一真实源，与主项目配置表惯例一致）：
  scan   扫描 MOD 项目 Assets/ModResource/Spine/BrownDust 资源目录(char_XXXXXX) → 重建道具 Excel（合并保留语言 Excel 的人工改名）
  export 读道具/语言 Excel → 导出 Mod JsonText（ItemsInfo.txt + Language_ItemsInfo_*.txt ×12）
  all    scan + export（默认）

与 AeonsEchoSpine/ArkReSpine/NikkeSpine 的差异：
  - 资源结构=目录(char_+6位数字 或 illust_*/npc* 等)/子目录(本体 或 本体_cutscene)/完整 spine 导出物；资源名=SkeletonData 资产文件名(去扩展名)
  - 目录两类编号（2026-09-28 与用户确认非 char 目录也出药）：char_XXXXXX 目录=数字部分(6位)进ID；非 char 目录（illust_dating/illust_special/illust_talk/npc/specialIllust/storypack 等,含 SkeletonData 才出药）=独立编号段 900001 起按目录自然序分配
  - 全部资源只有 default 一套皮肤（2026-09 全量 290 个 JSON 扫描确认），不按皮肤拆药
  - **出药规则=按动画名拆药**（2026-09-28 与用户确认）：
    动画名(小写)含 idle / all / loop / cut 的,**每个匹配动画各出 1 个幻化药**,该药的 ui_show_idle_anim 键=对应动画名;
    四类标签独立判定**不去重**——同一动画多重匹配时(如 1_cut_idle 同时含 idle+cut)在各类下分别出药;
    骨架内序号排序 = idle 系(动画名自然序) → all 系 → loop 系 → cut 系;
    0 匹配动画的骨架兜底出 1 个药(不带 idle_anim 键,静态显示)+警告（当前全量扫描不存在此类骨架）
  - 全部幻化药只改 ui_show（详情UI高清展示），不动 show/world（同 ArkRe/Nikke）

other_data 键值格式（& 拆项、每项首个 : 拆键值，同主项目 attack_mode other_data 规约，缺省键省略）：
  ui_show_res:000201_SkeletonData&ui_show_data:0.1919;0,0&ui_show_idle_anim:cut_A
  ui_show_res  ui_show_spine 高清展示资源名（详情UI，isUIShow=true 时使用）
  ui_show_data 详情UI尺寸「scale;x,y」（生成器按 645/骨架高 校准 scale，默认位移 0,0；scan 重建时默认按 id 保留手调值，--reset-layout 强制重算）
  ui_show_idle_anim 该药对应的目标动画名（按动画拆药规则必带；0 匹配兜底药省略=静态显示）

道具自ID规则：`18` + 套装号(char_XXXXXX 的 6 位数字；非 char 目录=900001 起独立编号段) + 2位序号(01 起)；name 自ID=道具自ID。
  套装内序号排序 = 骨架(子目录自然序:本体→cutscene) → 动画匹配(idle系→all系→loop系→cut系,系内动画名自然序)。

Excel 位置（MOD 项目，可用 Excel/WPS 直接打开查看/调整；与其他 Mod 的 Excel 相互独立）：
  Assets/Data/Excel/excel_mod_items_info_browndustspine[Mod道具信息-BrownDustSpine].xlsx - 道具配置（3 行表头，name[language] 标记列）
  Assets/Data/Excel/excel_mod_language_browndustspine[Mod多语言-BrownDustSpine].xlsx     - 道具名多语言（id + content_{12语言}）

注意：
  - scan 会全量重建道具表数据行（覆盖前先备份到 MOD项目/ExcelBackup/），但 **ui_show_data 手调值按 id 保留**（新增资源按骨架校准默认值；--reset-layout 可强制全部重算）；改了道具表参数后请只跑 export
  - 语言表按 id 合并保留人工内容：新增道具补默认名，删除的道具行自动清理
  - 规则详见主项目 .claude/skills/browndust-spine-mod/SKILL.md

用法（一律走主项目 .claude/scripts/run-python.ps1 包装）：
  powershell.exe -NoProfile -ExecutionPolicy Bypass -File ".claude/scripts/run-python.ps1" `
      ".claude/scripts/gen_browndust_spine_mod.py" all --mod-project "<MOD项目根>"
  # 构建完成后部署到主项目：
  ... export --mod-project "<MOD项目根>" --deploy-main "<主项目根>"
"""
import argparse
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl

MOD_NAME = "BrownDustSpine"
SOURCE_REL = "Assets/ModResource/Spine/BrownDust"
ICON_RES = "Item_TransformPotion_1"  # 复用主游戏内置幻化药图标
ITEM_TYPE_TRANSFORM_POTION = 18
LANGUAGES = ["cn", "en", "jp", "kr", "tw", "de", "fr", "ru", "es", "br", "pl", "tr"]

# 道具表：列名/类型/说明（布局与主项目 excel_items_info 一致，name 带 [language] 标记约定）
ITEM_SHEET = "ItemsInfo"
# 与主项目 TestTransformPotionGUI.GetModItemsExcelRelPath 的约定一致：excel_mod_items_info_{modName小写}[Mod道具信息-{modName}].xlsx
ITEM_EXCEL_REL = "Assets/Data/Excel/excel_mod_items_info_browndustspine[Mod道具信息-BrownDustSpine].xlsx"
ITEM_COLUMNS = [
    ("id", "long", "道具ID(key)=18+套装号(char的6位数字)+2位序号"),
    ("item_type", "int", "道具类型(18=幻化药)"),
    ("item_weapon_type", "int", "武器类型(仅武器有效)"),
    ("num_max", "int", "堆叠上限(1=不堆叠)"),
    ("creature_model_id", "long", "生物模组信息ID(0=通用)"),
    ("creature_model_info_id", "long", "生物模组详细信息ID"),
    ("icon_res", "string", "图标资源"),
    ("icon_rotate_z", "float", "图标旋转"),
    ("attack_mode_data", "string", "攻击模式数据(幻化药不用)"),
    ("other_data", "string", "形象键值串:ui_show_res:X&ui_show_data:scale;x,y[&ui_show_idle_anim:动画名](&拆项,:拆键值,缺省键省略;只改详情UI,不动show/world;idle_anim=按动画拆药的目标动画)"),
    ("name[language]", "long", "道具名textId(=道具id,文本在excel_mod_language_browndustspine)"),
    ("remark", "string", "备注"),
    ("reward_rarity", "string", "奖励稀有度白名单(空=全适配;消耗品不进装备池)"),
    ("source", "string", "道具来源(逗号分隔枚举值,空=默认来源) 1=征服模式奖励"),
]

# 语言表：id + content_{lang}（与主项目 excel_language 的 sheet 布局一致；与其他 Mod 的语言 Excel 相互独立）
LANG_SHEET = "ItemsInfo"
LANG_EXCEL_REL = "Assets/Data/Excel/excel_mod_language_browndustspine[Mod多语言-BrownDustSpine].xlsx"

# 资源目录名规则：char_+数字（如 char_000201）；数字部分进道具ID
SET_DIR_PATTERN = re.compile(r"^char_(\d+)$")
# 非 char 目录（illust_*/npc* 等）出药的套装号起始段：900001 起按目录自然序递增分配（2026-09-28 与用户确认；
# char 数字最大 099999<900000 不冲突；新增非 char 目录会自然顺延,已有目录编号不变=已生成 id 不漂移）
NON_CHAR_SET_NUM_BASE = 900000
# 出药动画标签及骨架内序号排序优先级：动画名(小写)含该子串即匹配出药；四类独立判定不去重(2026-09-28 与用户确认)
ANIM_TAGS = ("idle", "all", "loop", "cut")


def natural_key(name: str):
    """自然排序键（数字段按数值）：char_000201 < char_001101；000201 < 000201_cutscene"""
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", name)]


def default_name(lang: str, set_id: str, seq: int, anim: str) -> str:
    """道具名默认值：cn/tw 中文，其余语言英文兜底；按动画拆药的药带动画名后缀（0匹配兜底药无后缀）"""
    anim_tag = f" {anim}" if anim else ""
    if lang == "cn":
        return f"幻化药·棕尘{set_id}-{seq:02d}{anim_tag}"
    if lang == "tw":
        return f"幻化藥·棕塵{set_id}-{seq:02d}{anim_tag}"
    return f"BrownDust Potion {set_id}-{seq:02d}{anim_tag}"


def get_spine_height(json_path: Path) -> float:
    """读 spine json 的骨架包围盒高度（用于校准详情UI缩放）"""
    try:
        data = json.loads(json_path.read_text(encoding="utf-8"))
        return float(data["skeleton"]["height"])
    except Exception:
        return 0.0


def get_spine_anims(json_path: Path) -> list:
    """读 spine json 的动画名列表（4.x 格式 animations 为 dict,取 keys 保持声明顺序）；异常返回空列表"""
    try:
        data = json.loads(json_path.read_text(encoding="utf-8"))
        anims = data.get("animations")
        if isinstance(anims, dict):
            return list(anims.keys())
    except Exception:
        pass
    return []


def backup_excel(excel_path: Path, mod_project: Path):
    """覆盖 Excel 前备份到 MOD项目/ExcelBackup/（Assets 之外，不被导出工具扫描、不产生 .meta）；
    滚动复用 .bak.1~.bak.3（1=最新）——移位覆盖旧文件不新增，目录里永远只有最近 3 份，并顺带清掉旧版时间戳命名备份"""
    if not excel_path.exists():
        return
    backup_dir = mod_project / "ExcelBackup"
    backup_dir.mkdir(exist_ok=True)
    name = excel_path.name
    # 移位覆盖：bak.3←bak.2、bak.2←bak.1（replace 直接覆盖复用旧文件，不产生新文件）
    for i in range(3, 1, -1):
        newer = backup_dir / f"{name}.bak.{i - 1}"
        if newer.exists():
            newer.replace(backup_dir / f"{name}.bak.{i}")
    shutil.copy2(excel_path, backup_dir / f"{name}.bak.1")
    # 清理非滚动命名的旧备份（时间戳版等），保证永远都只有 3 份
    keep = {f"{name}.bak.{i}" for i in range(1, 4)}
    for old in backup_dir.iterdir():
        if old.is_file() and old.name.startswith(f"{name}.bak.") and old.name not in keep:
            old.unlink()


def build_other_data(ui_show_res: str, ui_show_data: str = "", ui_show_idle_anim: str = "") -> str:
    """拼 other_data 键值串：& 拆项、每项首个 : 拆键值（同主项目 attack_mode other_data 规约），缺省键省略；
    BrownDustSpine 药只改 ui_show（详情UI），不动 show/world；ui_show_idle_anim=按动画拆药的目标动画名（0匹配兜底药省略）"""
    segs = []
    if ui_show_res:
        segs.append(f"ui_show_res:{ui_show_res}")
    if ui_show_data:
        segs.append(f"ui_show_data:{ui_show_data}")
    if ui_show_idle_anim:
        segs.append(f"ui_show_idle_anim:{ui_show_idle_anim}")
    return "&".join(segs)


def read_preserved_ui_show_data(excel_path: Path) -> dict:
    """读现有道具表提取各行 other_data 的 ui_show_data 键 → {道具id: ui_show_data值}；
    scan 全量重建时按 id 保留测试面板手调的详情UI尺寸（同 id=同目录同资源同动画,骨架未变,手调值仍有效；新增按骨架校准）"""
    preserved = {}
    if not excel_path.exists():
        return preserved
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    if ITEM_SHEET not in wb.sheetnames:
        wb.close()
        return preserved
    ws = wb[ITEM_SHEET]
    header = [c.value for c in ws[1]]
    if "id" not in header or "other_data" not in header:
        wb.close()
        return preserved
    col_id = header.index("id")
    col_other = header.index("other_data")
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None or row[col_other] is None:
            continue
        for seg in str(row[col_other]).split("&"):
            if seg.startswith("ui_show_data:"):
                preserved[int(row[col_id])] = seg[len("ui_show_data:"):].strip()
                break
    wb.close()
    return preserved


def pick_tag_anims(anims: list):
    """按出药动画标签检出 (标签, 动画名) 列表：四类标签独立判定不去重（同一动画多重匹配在各类下分别出药），
    排序 = 标签优先级(idle→all→loop→cut) → 系内动画名自然序"""
    picked = []
    for tag in ANIM_TAGS:
        for a in sorted([x for x in anims if tag in x.lower()], key=natural_key):
            picked.append((tag, a))
    return picked


def compute_items(source_dir: Path, ui_scale_k: float, ui_pos_y: float, preserved_ui: dict = None):
    """扫描资源目录计算道具行，返回 (道具行列表, 跳过清单, 警告清单, 出药统计dict)；道具行 key 为干净字段名；
    preserved_ui=scan 前从旧道具表读出的 ui_show_data 手调值（按 id 保留）"""
    preserved_ui = preserved_ui or {}
    items = []
    used_ids = set()
    warnings = []
    skipped = []
    # 出药统计：四类标签各出药数 + 0匹配骨架数
    tag_stats = {tag: 0 for tag in ANIM_TAGS}
    no_match_count = 0
    # 非 char 目录（illust/npc 等）套装号分配计数（900001 起按目录自然序递增）
    non_char_seq = 0

    for set_dir in sorted(source_dir.iterdir(), key=lambda p: natural_key(p.name)):
        if not set_dir.is_dir():
            continue
        m = SET_DIR_PATTERN.match(set_dir.name)
        # 收集目录内全部骨架（递归子目录,子目录自然序=本体→cutscene）：每个 *_SkeletonData.asset
        # 资源名=资产文件名去扩展名(保留 _SkeletonData, 与构建器 Address=Path.GetFileNameWithoutExtension 及其他 Mod 约定一致)；
        # 同名 spine json=文件名去 _SkeletonData.asset + .json
        skeletons = []  # [(资源名, json路径)]
        for sd_file in sorted(set_dir.rglob("*_SkeletonData.asset"), key=lambda p: natural_key(str(p.relative_to(set_dir)))):
            res_name = sd_file.name[: -len(".asset")]
            json_path = sd_file.parent / f"{sd_file.name[: -len('_SkeletonData.asset')]}.json"
            if not json_path.exists():
                warnings.append(f"{set_dir.name}: {res_name} 缺同名 spine json,跳过该资源")
                continue
            skeletons.append((res_name, json_path))
        if not skeletons:
            skipped.append(f"{set_dir.name}(无 spine 资源)")
            continue
        if m:
            set_num = m.group(1)
        else:
            # 非 char 目录（illust/npc 等）也出药：独立编号段 900001 起
            non_char_seq += 1
            set_num = str(NON_CHAR_SET_NUM_BASE + non_char_seq)

        seq = 0
        for res_name, json_path in skeletons:
            anims = get_spine_anims(json_path)
            if not anims:
                warnings.append(f"{set_dir.name}: {res_name} 动画列表读取失败,跳过该资源")
                continue
            # 出药规则=按动画名拆药:idle/all/loop/cut 四类独立判定不去重,每个匹配动画各出 1 药
            picked = pick_tag_anims(anims)
            if not picked:
                # 0 匹配兜底:出 1 个无 idle_anim 键的药(静态显示)+警告（当前全量扫描不存在此类骨架）
                no_match_count += 1
                warnings.append(f"{set_dir.name}: {res_name} 无 idle/all/loop/cut 动画,兜底出 1 药(静态显示)")
                picked = [("", "")]
            # ui_show_data 按骨架校准一次(该骨架全部药共用,手调值按 id 各自保留)
            height = get_spine_height(json_path)
            for tag, anim in picked:
                seq += 1
                if seq > 99:
                    warnings.append(f"{set_dir.name}: 组合数超过99,ID序号溢出,后续组合未生成")
                    break
                self_id = int(f"{ITEM_TYPE_TRANSFORM_POTION}{set_num}{seq:02d}")
                if self_id in used_ids:
                    warnings.append(f"ID冲突: {self_id} ({set_dir.name} 第{seq}个组合)")
                    continue
                used_ids.add(self_id)
                if tag:
                    tag_stats[tag] += 1
                # ui_show_data: 手调值按 id 保留优先,否则按骨架高校准 scale=K/骨架高
                ui_show_data = preserved_ui.get(self_id, "")
                if not ui_show_data:
                    ui_scale = round(ui_scale_k / height, 4) if height > 0 else 0.12
                    ui_show_data = f"{ui_scale};0,{int(ui_pos_y)}"
                other_data = build_other_data(res_name, ui_show_data, anim)
                items.append({
                    "id": self_id,
                    "item_type": ITEM_TYPE_TRANSFORM_POTION,
                    "item_weapon_type": 0,
                    "num_max": 1,
                    "creature_model_id": 0,
                    "creature_model_info_id": 0,
                    "icon_res": ICON_RES,
                    "icon_rotate_z": 0.0,
                    "attack_mode_data": "",
                    "other_data": other_data,
                    "name": self_id,
                    "remark": f"棕尘幻化药(BrownDust {set_dir.name} {res_name}{(' ' + anim) if anim else ''})",
                    "reward_rarity": "",
                    # 来源=征服模式奖励(ItemSourceEnum.ConquerReward)：征服通关领奖的魔晶位之一随机替换为幻化药
                    "source": "1",
                    "_set_id": set_dir.name,  # 内部字段：供语言默认名使用，不进 Excel/JSON
                    "_seq": seq,
                    "_anim": anim,
                })
    return items, skipped, warnings, (tag_stats, no_match_count)


def write_item_excel(excel_path: Path, items, mod_project: Path):
    """全量重建道具表（3 行表头：列名/类型/说明；数据从第 4 行起）"""
    backup_excel(excel_path, mod_project)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ITEM_SHEET
    for col, (name, type_name, comment) in enumerate(ITEM_COLUMNS, start=1):
        ws.cell(row=1, column=col, value=name)
        ws.cell(row=2, column=col, value=type_name)
        ws.cell(row=3, column=col, value=comment)
    for row, item in enumerate(items, start=4):
        for col, (name, _, _) in enumerate(ITEM_COLUMNS, start=1):
            key = name.replace("[language]", "")
            ws.cell(row=row, column=col, value=item[key])
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)


def merge_language_excel(excel_path: Path, items, mod_project: Path):
    """语言表按 id 合并：保留人工改名，新道具补默认名，失效 id 清理"""
    existing = {}  # id -> {lang: content}
    if excel_path.exists():
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        if LANG_SHEET in wb.sheetnames:
            ws = wb[LANG_SHEET]
            header = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=4, values_only=True):
                if row[0] is None:
                    continue
                row_id = int(row[0])
                existing[row_id] = {}
                for i, col_name in enumerate(header):
                    if col_name and col_name.startswith("content_") and i < len(row):
                        lang = col_name.replace("content_", "")
                        if row[i] is not None and str(row[i]).strip() != "":
                            existing[row_id][lang] = str(row[i])
        wb.close()

    backup_excel(excel_path, mod_project)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = LANG_SHEET
    ws.cell(row=1, column=1, value="id")
    ws.cell(row=2, column=1, value="long")
    ws.cell(row=3, column=1, value="道具名textId(key)=道具id")
    for col, lang in enumerate(LANGUAGES, start=2):
        ws.cell(row=1, column=col, value=f"content_{lang}")
        ws.cell(row=2, column=col, value="string")
        ws.cell(row=3, column=col, value="道具名(空=英文兜底)")
    kept = 0
    for row, item in enumerate(items, start=4):
        self_id = item["id"]
        ws.cell(row=row, column=1, value=self_id)
        old = existing.get(self_id, {})
        for col, lang in enumerate(LANGUAGES, start=2):
            content = old.get(lang) or default_name(lang, item["_set_id"], item["_seq"], item["_anim"])
            if old.get(lang):
                kept += 1
            ws.cell(row=row, column=col, value=content)
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)
    return kept


def read_item_excel(excel_path: Path):
    """读道具表 → JSON 行（key 剥离 [language] 标记，按类型还原数值）"""
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb[ITEM_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = rows[0]
    types = rows[1]
    result = []
    for row in rows[3:]:
        if row[0] is None:
            continue
        item = {}
        for i, col_name in enumerate(header):
            if col_name is None:
                continue
            key = str(col_name).replace("[language]", "")
            type_name = str(types[i]) if types[i] else "string"
            value = row[i] if i < len(row) else None
            if type_name in ("long", "int"):
                item[key] = int(value) if value is not None and str(value).strip() != "" else 0
            elif type_name == "float":
                item[key] = float(value) if value is not None and str(value).strip() != "" else 0.0
            else:
                item[key] = "" if value is None else str(value)
        result.append(item)
    result.sort(key=lambda x: x["id"])
    return result


def read_language_excel(excel_path: Path):
    """读语言表 → {lang: [{id, content}]}（空内容按 cn→en 顺序兜底）"""
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb[LANG_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(c) if c else "" for c in rows[0]]
    per_lang = {lang: {} for lang in LANGUAGES}
    for row in rows[3:]:
        if row[0] is None:
            continue
        row_id = int(row[0])
        contents = {}
        for i, col_name in enumerate(header):
            if col_name.startswith("content_") and i < len(row) and row[i] is not None and str(row[i]).strip() != "":
                contents[col_name.replace("content_", "")] = str(row[i])
        for lang in LANGUAGES:
            # 兜底链：本语言 → cn → en → id
            content = contents.get(lang) or contents.get("cn") or contents.get("en") or str(row_id)
            per_lang[lang][row_id] = content
    return per_lang


def main():
    parser = argparse.ArgumentParser(description="BrownDustSpine Mod 数据生成器（Excel→JsonText 两段式）")
    parser.add_argument("command", nargs="?", default="all", choices=["scan", "export", "all"],
                        help="scan=资源→Excel；export=Excel→JsonText；all=两者（默认）")
    parser.add_argument("--mod-project", required=True, help="MOD 项目根目录")
    parser.add_argument("--deploy-main", default="", help="可选：主项目根目录，构建完成后把 Mod 整个目录部署过去")
    parser.add_argument("--ui-scale-k", type=float, default=645.0, help="详情UI缩放校准常数 K：scale=K/骨架高（默认645，与其他 Mod 同基准）")
    parser.add_argument("--ui-pos-y", type=float, default=0.0, help="详情UI默认Y偏移（默认0）")
    parser.add_argument("--reset-layout", action="store_true",
                        help="重建时重新设置卡片位置/大小（ui_show_data 全部按骨架重算）；默认关闭=按 id 保留已有手调值，仅新资源计算默认值")
    args = parser.parse_args()

    mod_project = Path(args.mod_project)
    source_dir = mod_project / SOURCE_REL
    item_excel = mod_project / ITEM_EXCEL_REL
    lang_excel = mod_project / LANG_EXCEL_REL

    if args.command in ("scan", "all"):
        if not source_dir.is_dir():
            print(f"[错误] 资源目录不存在: {source_dir}")
            sys.exit(1)
        preserved_ui = {} if args.reset_layout else read_preserved_ui_show_data(item_excel)
        items, skipped, warnings, (tag_stats, no_match_count) = compute_items(source_dir, args.ui_scale_k, args.ui_pos_y, preserved_ui)
        write_item_excel(item_excel, items, mod_project)
        kept = merge_language_excel(lang_excel, items, mod_project)
        print(f"[scan] 重建道具表 {len(items)} 行 → {item_excel}")
        print(f"[scan] 合并语言表（保留人工内容 {kept} 格）→ {lang_excel}")
        print(f"[scan] 按动画拆药统计：idle {tag_stats['idle']}，all {tag_stats['all']}，loop {tag_stats['loop']}，cut {tag_stats['cut']}，0匹配兜底 {no_match_count}（四类独立判定不去重）")
        if args.reset_layout:
            print("[scan] --reset-layout：ui_show_data 全部按骨架重算")
        elif preserved_ui:
            print(f"[scan] 按 id 保留手调 ui_show_data {len(preserved_ui)} 行（--reset-layout 可强制重算）")
        if skipped:
            print(f"[scan] 跳过: {', '.join(skipped)}")
        for w in warnings:
            print(f"[scan] [警告] {w}")

    if args.command in ("export", "all"):
        if not item_excel.exists() or not lang_excel.exists():
            print("[错误] Excel 不存在，请先运行 scan")
            sys.exit(1)
        items = read_item_excel(item_excel)
        per_lang = read_language_excel(lang_excel)

        out_dir = mod_project / "Mods" / MOD_NAME / "JsonText"
        out_dir.mkdir(parents=True, exist_ok=True)
        dump = lambda rows: json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
        (out_dir / "ItemsInfo.txt").write_text(dump(items), encoding="utf-8")
        for lang in LANGUAGES:
            rows = [{"id": k, "content": v} for k, v in sorted(per_lang[lang].items())]
            (out_dir / f"Language_ItemsInfo_{lang}.txt").write_text(dump(rows), encoding="utf-8")
        print(f"[export] 导出 {len(items)} 个道具 → {out_dir}")
        print(f"[export] ItemsInfo.txt + Language_ItemsInfo_*.txt ×{len(LANGUAGES)}")

    # 部署：构建完成后把 Mod 整个目录拷贝到主项目 Mods 下
    if args.deploy_main:
        mod_out = mod_project / "Mods" / MOD_NAME
        if not (mod_out / "catalog.bin").exists():
            print(f"[错误] 未找到构建产物 {mod_out}/catalog.bin —— 请先在 MOD 项目 Unity 里执行「工具/Mod/BrownDustSpine/一键构建」")
            sys.exit(1)
        target = Path(args.deploy_main) / "Mods" / MOD_NAME
        if target.exists():
            shutil.rmtree(target)
        shutil.copytree(mod_out, target)
        print(f"[deploy] 已部署 → {target}")
        print("下一步：主项目 Unity 里用 LauncherTest 启动（内含 InitializeAllModsSync），吃幻化药验证详情UI形象/动画")


if __name__ == "__main__":
    main()
